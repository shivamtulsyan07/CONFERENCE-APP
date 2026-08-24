"""Backend tests for iteration 14:

  * GET /api/balance-stock  — balance = stock + pending(sent-arrived, >0) - ordered - shop_sold
  * GET /api/extra-stock    — only lines where arrived > sent, plus totals
  * GET /api/extra-stock/export.xlsx  — valid xlsx with header + TOTAL row
  * GET /api/extra-stock/export.pdf   — valid pdf

All rows use a QA_XS_ prefix and are cleaned up after each test.
"""
import io
import os
import pytest
import requests
from pathlib import Path
from dotenv import load_dotenv

load_dotenv(Path(__file__).resolve().parents[2] / "frontend" / ".env")
BASE_URL = os.environ["REACT_APP_BACKEND_URL"].rstrip("/")
API = f"{BASE_URL}/api"


@pytest.fixture(scope="module")
def s():
    sess = requests.Session()
    sess.headers.update({"Content-Type": "application/json"})
    return sess


def _cleanup_by_item(s, item_marker):
    """Delete every QA_XS_ row we may have created for `item_marker` across all
    line sheets, orders and stock. Best effort — never raises."""
    for sheet in ("company-sent-rows", "company-arrived-rows", "shop-sale-rows"):
        try:
            rows = s.get(f"{API}/line-sheet/{sheet}").json()
            for r in rows:
                if item_marker and item_marker.upper() in str(r.get("item", "")).upper():
                    s.delete(f"{API}/line-sheet/{sheet}/{r['id']}")
        except Exception:
            pass
    try:
        for r in s.get(f"{API}/order-rows").json():
            if item_marker.upper() in str(r.get("item", "")).upper():
                s.delete(f"{API}/order-rows/{r['id']}")
    except Exception:
        pass
    try:
        for r in s.get(f"{API}/stock-rows").json():
            if item_marker.upper() in str(r.get("item", "")).upper():
                s.delete(f"{API}/stock-rows/{r['id']}")
    except Exception:
        pass


# ---------- Balance-Stock formula ----------
class TestBalanceStockFormula:
    """balance = stock_qty + pending_with_company (max(sent-arrived,0)) - ordered_qty - shop_sold_qty"""

    def test_full_formula_with_pending_and_shop_sales(self, s):
        marker = "QA_XS_BAL_ITEM_ZZ9"
        _cleanup_by_item(s, marker)
        created = {"stock": None, "order": None, "sent": None, "arrived": None, "shop": None}
        try:
            group, shade = "QA_XS_G", "SS9"
            created["stock"] = s.post(f"{API}/stock-rows/bulk", json={"rows": [
                {"group": group, "item": marker, "shade": shade, "quantity": 20, "row_index": 700}
            ]}).json()["saved"][0]["id"]
            created["order"] = s.post(f"{API}/order-rows/bulk", json={"rows": [
                {"party_name": "QA_XS_P", "group": group, "item": marker,
                 "shade": shade, "qty": 12, "rate": 0, "bill_no": ""}
            ]}).json()["saved"][0]["id"]
            created["sent"] = s.post(f"{API}/line-sheet/company-sent-rows/bulk", json={"rows": [
                {"group": group, "item": marker, "shade": shade, "quantity": 30, "row_index": 700}
            ]}).json()["saved"][0]["id"]
            created["arrived"] = s.post(f"{API}/line-sheet/company-arrived-rows/bulk", json={"rows": [
                {"group": group, "item": marker, "shade": shade, "quantity": 10, "row_index": 700}
            ]}).json()["saved"][0]["id"]
            created["shop"] = s.post(f"{API}/line-sheet/shop-sale-rows/bulk", json={"rows": [
                {"party_name": "QA_SHOPP", "bill_no": "QA/B1",
                 "group": group, "item": marker, "shade": shade, "quantity": 3, "row_index": 700}
            ]}).json()["saved"][0]["id"]

            d = s.get(f"{API}/balance-stock").json()
            row = [r for r in d["rows"] if r["item"].upper() == marker][0]
            # stock=20, pending=30-10=20, ordered=12, shop=3 -> balance = 25
            assert row["stock_qty"] == 20
            assert row["ordered_qty"] == 12
            assert row["shop_sold_qty"] == 3
            assert row["balance"] == 25, f"expected 25 got {row['balance']} in {row}"

            # totals reconcile
            assert abs(d["total_stock"] - sum(r["stock_qty"] for r in d["rows"])) < 1e-6
            assert abs(d["total_ordered"] - sum(r["ordered_qty"] for r in d["rows"])) < 1e-6
            assert abs(d["total_balance"] - sum(r["balance"] for r in d["rows"])) < 1e-6
            assert d["short_lines"] == sum(1 for r in d["rows"] if r["balance"] < 0)
        finally:
            if created["order"]:   s.delete(f"{API}/order-rows/{created['order']}")
            if created["stock"]:   s.delete(f"{API}/stock-rows/{created['stock']}")
            if created["sent"]:    s.delete(f"{API}/line-sheet/company-sent-rows/{created['sent']}")
            if created["arrived"]: s.delete(f"{API}/line-sheet/company-arrived-rows/{created['arrived']}")
            if created["shop"]:    s.delete(f"{API}/line-sheet/shop-sale-rows/{created['shop']}")
            _cleanup_by_item(s, marker)

    def test_pending_never_negative(self, s):
        """When arrived >= sent, pending contribution to balance must be 0 (never negative)."""
        marker = "QA_XS_BAL_NONEG_88"
        _cleanup_by_item(s, marker)
        ids = {}
        try:
            group, shade = "QA_XS_G2", "N1"
            ids["stock"] = s.post(f"{API}/stock-rows/bulk", json={"rows": [
                {"group": group, "item": marker, "shade": shade, "quantity": 5, "row_index": 701}
            ]}).json()["saved"][0]["id"]
            ids["sent"] = s.post(f"{API}/line-sheet/company-sent-rows/bulk", json={"rows": [
                {"group": group, "item": marker, "shade": shade, "quantity": 4, "row_index": 701}
            ]}).json()["saved"][0]["id"]
            ids["arrived"] = s.post(f"{API}/line-sheet/company-arrived-rows/bulk", json={"rows": [
                {"group": group, "item": marker, "shade": shade, "quantity": 10, "row_index": 701}
            ]}).json()["saved"][0]["id"]
            d = s.get(f"{API}/balance-stock").json()
            row = [r for r in d["rows"] if r["item"].upper() == marker][0]
            # pending clamps to 0: balance = 5 + 0 - 0 - 0 = 5
            assert row["balance"] == 5, row
        finally:
            for k, v in ids.items():
                if k == "stock":
                    s.delete(f"{API}/stock-rows/{v}")
                elif k == "sent":
                    s.delete(f"{API}/line-sheet/company-sent-rows/{v}")
                elif k == "arrived":
                    s.delete(f"{API}/line-sheet/company-arrived-rows/{v}")
            _cleanup_by_item(s, marker)


# ---------- Extra Stock ----------
class TestExtraStock:
    def test_shape_and_totals_reconcile(self, s):
        r = s.get(f"{API}/extra-stock")
        assert r.status_code == 200
        d = r.json()
        for k in ["rows", "total_lines", "total_sent", "total_arrived", "total_extra"]:
            assert k in d
        assert d["total_lines"] == len(d["rows"])
        assert abs(d["total_sent"] - sum(r["sent_qty"] for r in d["rows"])) < 1e-6
        assert abs(d["total_arrived"] - sum(r["arrived_qty"] for r in d["rows"])) < 1e-6
        assert abs(d["total_extra"] - sum(r["extra_qty"] for r in d["rows"])) < 1e-6
        for row in d["rows"]:
            assert row["extra_qty"] > 0
            assert row["arrived_qty"] > row["sent_qty"]
            for k in ["group", "item", "shade", "sent_qty", "arrived_qty", "extra_qty", "last_date"]:
                assert k in row

    def test_only_lines_where_arrived_gt_sent_appear(self, s):
        marker_extra  = "QA_XS_EXTRA_ITEM_A"
        marker_equal  = "QA_XS_EQUAL_ITEM_B"
        marker_short  = "QA_XS_SHORT_ITEM_C"
        marker_blank  = "QA_XS_BLANK_ITEM_D"  # arrived row with no item — must be ignored
        for m in (marker_extra, marker_equal, marker_short, marker_blank):
            _cleanup_by_item(s, m)

        ids = []
        try:
            # EXTRA line: sent 5, arrived 8 -> extra 3
            ids.append(("company-sent-rows", s.post(f"{API}/line-sheet/company-sent-rows/bulk", json={"rows": [
                {"group": "GX", "item": marker_extra, "shade": "1", "quantity": 5,
                 "date": "2026-01-01", "row_index": 750}
            ]}).json()["saved"][0]["id"]))
            ids.append(("company-arrived-rows", s.post(f"{API}/line-sheet/company-arrived-rows/bulk", json={"rows": [
                {"group": "GX", "item": marker_extra, "shade": "1", "quantity": 8,
                 "date": "2026-01-05", "row_index": 750}
            ]}).json()["saved"][0]["id"]))
            # EQUAL line: sent 4 arrived 4 -> not in extra
            ids.append(("company-sent-rows", s.post(f"{API}/line-sheet/company-sent-rows/bulk", json={"rows": [
                {"group": "GX", "item": marker_equal, "shade": "1", "quantity": 4, "row_index": 751}
            ]}).json()["saved"][0]["id"]))
            ids.append(("company-arrived-rows", s.post(f"{API}/line-sheet/company-arrived-rows/bulk", json={"rows": [
                {"group": "GX", "item": marker_equal, "shade": "1", "quantity": 4, "row_index": 751}
            ]}).json()["saved"][0]["id"]))
            # SHORT line: sent 10 arrived 2 -> not in extra
            ids.append(("company-sent-rows", s.post(f"{API}/line-sheet/company-sent-rows/bulk", json={"rows": [
                {"group": "GX", "item": marker_short, "shade": "1", "quantity": 10, "row_index": 752}
            ]}).json()["saved"][0]["id"]))
            ids.append(("company-arrived-rows", s.post(f"{API}/line-sheet/company-arrived-rows/bulk", json={"rows": [
                {"group": "GX", "item": marker_short, "shade": "1", "quantity": 2, "row_index": 752}
            ]}).json()["saved"][0]["id"]))
            # BLANK item: arrived only with no item -> must be ignored (party_name+bill_no keeps it persistable)
            ids.append(("company-arrived-rows", s.post(f"{API}/line-sheet/company-arrived-rows/bulk", json={"rows": [
                {"party_name": "QA_XS_BLANKP", "bill_no": "QA/BLK",
                 "group": "GX", "item": "", "shade": "", "quantity": 999, "row_index": 753}
            ]}).json()["saved"][0]["id"]))

            d = s.get(f"{API}/extra-stock").json()
            items = {r["item"].upper() for r in d["rows"]}
            assert marker_extra in items
            assert marker_equal not in items
            assert marker_short not in items
            # extra row values
            xrow = [r for r in d["rows"] if r["item"].upper() == marker_extra][0]
            assert xrow["sent_qty"] == 5
            assert xrow["arrived_qty"] == 8
            assert xrow["extra_qty"] == 3
            assert xrow["last_date"] == "2026-01-05"

            # blank-item arrived row must not create a phantom entry
            for r in d["rows"]:
                assert str(r["item"]).strip() != ""
        finally:
            for sheet, rid in ids:
                s.delete(f"{API}/line-sheet/{sheet}/{rid}")
            for m in (marker_extra, marker_equal, marker_short, marker_blank):
                _cleanup_by_item(s, m)


# ---------- Extra Stock exports ----------
class TestExtraStockExports:
    def test_xlsx_headers_and_total_row(self, s):
        r = s.get(f"{API}/extra-stock/export.xlsx")
        assert r.status_code == 200
        assert "spreadsheetml" in r.headers.get("content-type", "")
        assert len(r.content) > 500
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        header = [ws.cell(row=4, column=c).value for c in range(1, 8)]
        assert header == ["SR", "GROUP NAME", "ITEM NAME", "SHADE",
                          "SENT QTY", "ARRIVED QTY", "EXTRA QTY"]
        api_rows = s.get(f"{API}/extra-stock").json()["rows"]
        assert ws.max_row == 4 + len(api_rows) + 1
        assert str(ws.cell(row=ws.max_row, column=4).value).strip().upper() == "TOTAL"
        assert ws.cell(row=ws.max_row, column=4).font.bold is True
        assert ws.cell(row=ws.max_row, column=7).value == sum(x["extra_qty"] for x in api_rows)

    def test_pdf_valid(self, s):
        r = s.get(f"{API}/extra-stock/export.pdf")
        assert r.status_code == 200
        assert r.headers.get("content-type", "").startswith("application/pdf")
        assert r.content[:4] == b"%PDF"
        assert len(r.content) > 500


# ---------- Regression: all critical endpoints still 200 ----------
class TestRegressionEndpoints:
    @pytest.mark.parametrize("path", [
        "/parties", "/order-rows", "/stock-rows", "/lookups",
        "/company-balance", "/company-order?pending_only=true",
        "/company-order?pending_only=false", "/balance-stock",
        "/order-summary", "/extra-stock", "/stats/dashboard", "/stats/overview",
        "/line-sheet/company-sent-rows", "/line-sheet/company-arrived-rows",
        "/line-sheet/shop-sale-rows",
    ])
    def test_endpoint_ok(self, s, path):
        r = s.get(f"{API}{path}")
        assert r.status_code == 200, f"{path} -> {r.status_code} {r.text[:200]}"
