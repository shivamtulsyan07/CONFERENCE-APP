from fastapi import FastAPI, APIRouter, HTTPException
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
from starlette.responses import StreamingResponse
import io
import os
import logging
from pathlib import Path
from pydantic import BaseModel, Field, ConfigDict
from typing import List, Optional, Literal
from bson import ObjectId
from datetime import datetime, timezone

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

app = FastAPI()
api_router = APIRouter(prefix="/api")


def now_iso():
    return datetime.now(timezone.utc).isoformat()


def oid(value: str):
    if not ObjectId.is_valid(value):
        raise HTTPException(status_code=400, detail="Invalid id")
    return ObjectId(value)


class BaseDocument(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: Optional[str] = Field(default=None)

    def to_mongo(self):
        doc = self.model_dump(exclude_none=True)
        doc.pop("id", None)
        return doc

    @classmethod
    def from_mongo(cls, doc):
        if doc is None:
            return None
        d = dict(doc)
        if "_id" in d:
            d["id"] = str(d.pop("_id"))
        return cls(**d)


STATUS = Literal["not_ready", "arrived", "ready"]


# ---------- Parties ----------
class Party(BaseDocument):
    name: str
    page: str = ""
    phone: str = ""
    city: str = ""
    created_at: str = Field(default_factory=now_iso)


class PartyCreate(BaseModel):
    name: str
    page: str = ""
    phone: str = ""
    city: str = ""


# ---------- Order sheet row ----------
class OrderRow(BaseDocument):
    party_name: str = ""
    page: str = ""
    conference: str = ""
    group: str = ""
    item: str = ""
    shade: str = ""
    qty: float = 0
    mtr: str = ""
    rate: float = 0
    amount: float = 0
    bill_no: str = ""
    status: STATUS = "not_ready"
    row_index: int = 0
    created_at: str = Field(default_factory=now_iso)


class OrderRowIn(BaseModel):
    id: Optional[str] = None
    party_name: str = ""
    page: str = ""
    conference: str = ""
    group: str = ""
    item: str = ""
    shade: str = ""
    qty: float = 0
    mtr: str = ""
    rate: float = 0
    amount: float = 0
    bill_no: str = ""
    status: STATUS = "not_ready"
    row_index: int = 0


# ---------- Stock sheet row ----------
class StockRow(BaseDocument):
    group: str = ""
    item: str = ""
    shade: str = ""
    quantity: float = 0
    row_index: int = 0
    created_at: str = Field(default_factory=now_iso)


class StockRowIn(BaseModel):
    id: Optional[str] = None
    group: str = ""
    item: str = ""
    shade: str = ""
    quantity: float = 0
    row_index: int = 0


class LineRow(BaseDocument):
    group: str = ""
    item: str = ""
    shade: str = ""
    quantity: float = 0
    date: str = ""
    remark: str = ""
    row_index: int = 0
    created_at: str = Field(default_factory=now_iso)


class LineRowIn(BaseModel):
    id: Optional[str] = None
    group: str = ""
    item: str = ""
    shade: str = ""
    quantity: float = 0
    date: str = ""
    remark: str = ""
    row_index: int = 0


class BulkLineRows(BaseModel):
    rows: List[LineRowIn]


def is_blank_line(r: LineRowIn):
    return not any([r.item.strip(), r.shade.strip(), r.quantity, r.remark.strip()])


class BulkOrderRows(BaseModel):
    rows: List[OrderRowIn]


class BulkStockRows(BaseModel):
    rows: List[StockRowIn]


def is_blank_order(r: OrderRowIn):
    # conference / page / group alone are not real content
    return not any([
        r.party_name.strip(), r.item.strip(), r.shade.strip(), r.bill_no.strip(),
        r.mtr.strip(), r.qty, r.rate,
    ])


def is_blank_stock(r: StockRowIn):
    return not any([r.item.strip(), r.shade.strip(), r.quantity])


# ---------- Parties API ----------
@api_router.get("/parties", response_model=List[Party])
async def list_parties():
    docs = await db.parties.find().sort("name", 1).to_list(2000)
    return [Party.from_mongo(d) for d in docs]


@api_router.post("/parties", response_model=Party)
async def create_party(payload: PartyCreate):
    party = Party(**payload.model_dump())
    res = await db.parties.insert_one(party.to_mongo())
    party.id = str(res.inserted_id)
    return party


@api_router.delete("/parties/{party_id}")
async def delete_party(party_id: str):
    await db.parties.delete_one({"_id": oid(party_id)})
    return {"ok": True}


# ---------- Order sheet API ----------
@api_router.get("/order-rows", response_model=List[OrderRow])
async def list_order_rows():
    docs = await db.order_rows.find().sort([("row_index", 1), ("created_at", 1)]).to_list(5000)
    return [OrderRow.from_mongo(d) for d in docs]


@api_router.post("/order-rows/bulk")
async def save_order_rows(payload: BulkOrderRows):
    saved = []
    for i, r in enumerate(payload.rows):
        if is_blank_order(r):
            if r.id:
                await db.order_rows.delete_one({"_id": oid(r.id)})
            continue
        data = r.model_dump(exclude={"id"})
        data["amount"] = (r.qty or 0) * (r.rate or 0)
        if r.id:
            await db.order_rows.update_one({"_id": oid(r.id)}, {"$set": data})
            saved.append({"index": i, "id": r.id})
        else:
            data["created_at"] = now_iso()
            res = await db.order_rows.insert_one(data)
            saved.append({"index": i, "id": str(res.inserted_id)})
    return {"saved": saved}


@api_router.post("/order-rows/replace")
async def replace_order_rows(payload: BulkOrderRows):
    docs = []
    for i, r in enumerate(payload.rows):
        if is_blank_order(r):
            continue
        data = r.model_dump(exclude={"id"})
        data["amount"] = (r.qty or 0) * (r.rate or 0)
        data["row_index"] = i
        data["created_at"] = now_iso()
        docs.append(data)
    await db.order_rows.delete_many({})
    if docs:
        await db.order_rows.insert_many(docs)
    return {"count": len(docs)}


@api_router.delete("/order-rows/{row_id}")
async def delete_order_row(row_id: str):
    await db.order_rows.delete_one({"_id": oid(row_id)})
    return {"ok": True}


@api_router.post("/order-rows/auto-status", response_model=List[OrderRow])
async def auto_status():
    stock = await db.stock_rows.find().to_list(5000)
    available = {}
    for s in stock:
        key = (
            str(s.get("group", "")).strip().upper(),
            s.get("item", "").strip().upper(),
            str(s.get("shade", "")).strip(),
        )
        available[key] = available.get(key, 0) + (s.get("quantity") or 0)

    updated = 0
    rows = await db.order_rows.find().to_list(5000)
    for r in rows:
        key = (
            str(r.get("group", "")).strip().upper(),
            r.get("item", "").strip().upper(),
            str(r.get("shade", "")).strip(),
        )
        have = available.get(key, 0)
        need = r.get("qty") or 0
        if have <= 0:
            status = "not_ready"
        elif have >= need:
            status = "ready"
        else:
            status = "arrived"
        if status != r.get("status"):
            await db.order_rows.update_one({"_id": r["_id"]}, {"$set": {"status": status}})
            updated += 1
    return await list_order_rows()


# ---------- Stock sheet API ----------
@api_router.get("/stock-rows", response_model=List[StockRow])
async def list_stock_rows():
    docs = await db.stock_rows.find().sort([("row_index", 1), ("created_at", 1)]).to_list(5000)
    return [StockRow.from_mongo(d) for d in docs]


@api_router.post("/stock-rows/bulk", response_model=None)
@api_router.post("/stock-rows/bulk")
async def save_stock_rows(payload: BulkStockRows):
    saved = []
    for i, r in enumerate(payload.rows):
        if is_blank_stock(r):
            if r.id:
                await db.stock_rows.delete_one({"_id": oid(r.id)})
            continue
        data = r.model_dump(exclude={"id"})
        if r.id:
            await db.stock_rows.update_one({"_id": oid(r.id)}, {"$set": data})
            saved.append({"index": i, "id": r.id})
        else:
            data["created_at"] = now_iso()
            res = await db.stock_rows.insert_one(data)
            saved.append({"index": i, "id": str(res.inserted_id)})
    return {"saved": saved}


@api_router.post("/stock-rows/replace")
async def replace_stock_rows(payload: BulkStockRows):
    docs = []
    for i, r in enumerate(payload.rows):
        if is_blank_stock(r):
            continue
        data = r.model_dump(exclude={"id"})
        data["row_index"] = i
        data["created_at"] = now_iso()
        docs.append(data)
    await db.stock_rows.delete_many({})
    if docs:
        await db.stock_rows.insert_many(docs)
    return {"count": len(docs)}


@api_router.delete("/stock-rows/{row_id}")
async def delete_stock_row(row_id: str):
    await db.stock_rows.delete_one({"_id": oid(row_id)})
    return {"ok": True}


# ---------- Lookups ----------
@api_router.get("/lookups")
async def lookups():
    parties = await db.parties.find().to_list(2000)
    order_rows = await db.order_rows.find().to_list(5000)
    stock_rows = await db.stock_rows.find().to_list(5000)

    def uniq(values):
        return sorted({str(v).strip() for v in values if str(v).strip()})

    return {
        "parties": uniq([p["name"] for p in parties] + [r.get("party_name", "") for r in order_rows]),
        "party_pages": {
            p["name"]: p.get("page", "") for p in parties if p.get("page")
        },
        "conferences": uniq([r.get("conference", "") for r in order_rows] + [r.get("conference", "") for r in stock_rows]),
        "groups": uniq([r.get("group", "") for r in order_rows] + [r.get("group", "") for r in stock_rows] + ["SH ROLL"]),
        "items": uniq([r.get("item", "") for r in order_rows] + [r.get("item", "") for r in stock_rows]),
        "shades": uniq([r.get("shade", "") for r in order_rows] + [r.get("shade", "") for r in stock_rows]),
        "bill_nos": uniq([r.get("bill_no", "") for r in order_rows]),
    }


# ---------- Company sent / arrived line sheets ----------
LINE_COLLECTIONS = {
    "company-sent-rows": "company_sent_rows",
    "company-arrived-rows": "company_arrived_rows",
}


def line_collection(name: str):
    coll = LINE_COLLECTIONS.get(name)
    if not coll:
        raise HTTPException(status_code=404, detail="Unknown sheet")
    return db[coll]


@api_router.get("/line-sheet/{sheet}", response_model=List[LineRow])
async def list_line_rows(sheet: str):
    coll = line_collection(sheet)
    docs = await coll.find().sort([("row_index", 1), ("created_at", 1)]).to_list(10000)
    return [LineRow.from_mongo(d) for d in docs]


@api_router.post("/line-sheet/{sheet}/bulk")
async def save_line_rows(sheet: str, payload: BulkLineRows):
    coll = line_collection(sheet)
    saved = []
    for i, r in enumerate(payload.rows):
        if is_blank_line(r):
            if r.id:
                await coll.delete_one({"_id": oid(r.id)})
            continue
        data = r.model_dump(exclude={"id"})
        if r.id:
            await coll.update_one({"_id": oid(r.id)}, {"$set": data})
            saved.append({"index": i, "id": r.id})
        else:
            data["created_at"] = now_iso()
            res = await coll.insert_one(data)
            saved.append({"index": i, "id": str(res.inserted_id)})
    return {"saved": saved}


@api_router.post("/line-sheet/{sheet}/replace")
async def replace_line_rows(sheet: str, payload: BulkLineRows):
    coll = line_collection(sheet)
    docs = []
    for i, r in enumerate(payload.rows):
        if is_blank_line(r):
            continue
        data = r.model_dump(exclude={"id"})
        data["row_index"] = i
        data["created_at"] = now_iso()
        docs.append(data)
    await coll.delete_many({})
    if docs:
        await coll.insert_many(docs)
    return {"count": len(docs)}


@api_router.delete("/line-sheet/{sheet}/{row_id}")
async def delete_line_row(sheet: str, row_id: str):
    coll = line_collection(sheet)
    await coll.delete_one({"_id": oid(row_id)})
    return {"ok": True}


@api_router.get("/company-balance")
async def company_balance():
    sent = await db.company_sent_rows.find().to_list(20000)
    arrived = await db.company_arrived_rows.find().to_list(20000)

    agg = {}

    def key_of(r):
        return (
            str(r.get("group", "")).strip().upper(),
            str(r.get("item", "")).strip().upper(),
            str(r.get("shade", "")).strip(),
        )

    for r in sent:
        if not str(r.get("item", "")).strip():
            continue
        e = agg.setdefault(key_of(r), {
            "group": str(r.get("group", "")).strip(), "item": str(r.get("item", "")).strip(),
            "shade": str(r.get("shade", "")).strip(), "sent_qty": 0, "arrived_qty": 0,
        })
        e["sent_qty"] += r.get("quantity") or 0

    for r in arrived:
        if not str(r.get("item", "")).strip():
            continue
        e = agg.setdefault(key_of(r), {
            "group": str(r.get("group", "")).strip(), "item": str(r.get("item", "")).strip(),
            "shade": str(r.get("shade", "")).strip(), "sent_qty": 0, "arrived_qty": 0,
        })
        e["arrived_qty"] += r.get("quantity") or 0

    rows = []
    for e in agg.values():
        rows.append({**e, "balance_qty": e["sent_qty"] - e["arrived_qty"]})
    rows.sort(key=lambda x: (x["group"], x["item"], x["shade"]))
    return {
        "rows": rows,
        "total_sent": sum(r["sent_qty"] for r in rows),
        "total_arrived": sum(r["arrived_qty"] for r in rows),
        "total_balance": sum(r["balance_qty"] for r in rows),
        "pending_lines": sum(1 for r in rows if r["balance_qty"] > 0),
    }


# ---------- Company order (what to order from the company) ----------
def company_order_rows_sync(orders, stock):
    stock_by_key = {}
    for s in stock:
        key = (str(s.get("group", "")).strip().upper(), str(s.get("item", "")).strip().upper(), str(s.get("shade", "")).strip())
        stock_by_key[key] = stock_by_key.get(key, 0) + (s.get("quantity") or 0)

    agg = {}
    for o in orders:
        item = str(o.get("item", "")).strip()
        if not item:
            continue
        group = str(o.get("group", "")).strip()
        shade = str(o.get("shade", "")).strip()
        key = (group.upper(), item.upper(), shade)
        e = agg.setdefault(key, {
            "group": group, "item": item, "shade": shade,
            "ordered_qty": 0, "stock_qty": stock_by_key.get(key, 0), "parties": set(),
        })
        e["ordered_qty"] += o.get("qty") or 0
        party = str(o.get("party_name", "")).strip()
        if party:
            e["parties"].add(party)

    out = []
    for e in agg.values():
        to_order = e["ordered_qty"] - e["stock_qty"]
        out.append({
            "group": e["group"], "item": e["item"], "shade": e["shade"],
            "ordered_qty": e["ordered_qty"], "stock_qty": e["stock_qty"],
            "to_order": to_order if to_order > 0 else 0,
            "parties": len(e["parties"]),
        })
    out.sort(key=lambda x: (x["group"], x["item"], x["shade"]))
    return out


async def get_company_order(pending_only: bool):
    orders = await db.order_rows.find().to_list(20000)
    stock = await db.stock_rows.find().to_list(20000)
    rows = company_order_rows_sync(orders, stock)
    if pending_only:
        rows = [r for r in rows if r["to_order"] > 0]
    return rows


@api_router.get("/company-order")
async def company_order(pending_only: bool = False):
    rows = await get_company_order(pending_only)
    return {
        "rows": rows,
        "total_lines": len(rows),
        "total_ordered": sum(r["ordered_qty"] for r in rows),
        "total_stock": sum(r["stock_qty"] for r in rows),
        "total_to_order": sum(r["to_order"] for r in rows),
    }


@api_router.get("/company-order/export.xlsx")
async def company_order_xlsx(pending_only: bool = True):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    rows = await get_company_order(pending_only)
    wb = Workbook()
    ws = wb.active
    ws.title = "Company Order"
    ws.append(["COMPANY ORDER REQUIREMENT"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Generated {datetime.now(timezone.utc).strftime('%d-%m-%Y %H:%M UTC')}"])
    ws.append([])
    headers = ["SR", "GROUP NAME", "ITEM NAME", "SHADE", "ORDERED QTY", "IN HOUSE STOCK", "QTY TO ORDER"]
    ws.append(headers)
    head_row = ws.max_row
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=head_row, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0A2540")
        cell.alignment = Alignment(horizontal="center")
    for i, r in enumerate(rows, start=1):
        ws.append([i, r["group"], r["item"], r["shade"], r["ordered_qty"], r["stock_qty"], r["to_order"]])
    total_row = ws.max_row + 1
    ws.cell(row=total_row, column=4, value="TOTAL").font = Font(bold=True)
    for col, key in ((5, "ordered_qty"), (6, "stock_qty"), (7, "to_order")):
        ws.cell(row=total_row, column=col, value=sum(r[key] for r in rows)).font = Font(bold=True)
    for col, width in zip("ABCDEFG", (6, 22, 32, 12, 14, 16, 14)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = ws.cell(row=head_row + 1, column=1)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="company-order.xlsx"'},
    )


@api_router.get("/company-order/export.pdf")
async def company_order_pdf(pending_only: bool = True):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

    rows = await get_company_order(pending_only)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=15 * mm, bottomMargin=15 * mm,
                            leftMargin=12 * mm, rightMargin=12 * mm, title="Company Order")
    styles = getSampleStyleSheet()
    story = [
        Paragraph("<b>COMPANY ORDER REQUIREMENT</b>", styles["Title"]),
        Paragraph(f"Generated {datetime.now(timezone.utc).strftime('%d-%m-%Y %H:%M UTC')}", styles["Normal"]),
        Spacer(1, 6 * mm),
    ]
    data = [["SR", "GROUP NAME", "ITEM NAME", "SHADE", "ORDERED", "IN HOUSE", "TO ORDER"]]
    for i, r in enumerate(rows, start=1):
        data.append([str(i), r["group"], r["item"], r["shade"],
                     f'{r["ordered_qty"]:g}', f'{r["stock_qty"]:g}', f'{r["to_order"]:g}'])
    data.append(["", "", "", "TOTAL",
                 f'{sum(r["ordered_qty"] for r in rows):g}',
                 f'{sum(r["stock_qty"] for r in rows):g}',
                 f'{sum(r["to_order"] for r in rows):g}'])
    table = Table(data, repeatRows=1, colWidths=[12 * mm, 38 * mm, 58 * mm, 20 * mm, 20 * mm, 22 * mm, 22 * mm])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0A2540")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ALIGN", (4, 0), (-1, -1), "RIGHT"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C9D3E0")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#F4F7FA")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    story.append(table)
    doc.build(story)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/pdf",
        headers={"Content-Disposition": 'attachment; filename="company-order.pdf"'},
    )


# ---------- Conference order summary ----------
@api_router.get("/order-summary")
async def order_summary():
    orders = await db.order_rows.find().to_list(5000)
    agg = {}
    for o in orders:
        item = str(o.get("item", "")).strip()
        if not item:
            continue
        group = str(o.get("group", "")).strip()
        shade = str(o.get("shade", "")).strip()
        key = (group.upper(), item.upper(), shade)
        e = agg.setdefault(key, {"group": group, "item": item, "shade": shade, "quantity": 0, "rows": 0})
        e["quantity"] += o.get("qty") or 0
        e["rows"] += 1
    out = sorted(agg.values(), key=lambda x: (x["group"], x["item"], x["shade"]))
    return {
        "rows": out,
        "total_quantity": sum(r["quantity"] for r in out),
        "total_lines": len(out),
    }


# ---------- Balance stock ----------
@api_router.get("/balance-stock")
async def balance_stock():
    stock = await db.stock_rows.find().to_list(5000)
    orders = await db.order_rows.find().to_list(5000)

    rows = {}

    def entry(item, shade, group=""):
        key = (str(group).strip().upper(), item.strip().upper(), str(shade).strip())
        return rows.setdefault(key, {
            "group": str(group).strip(), "item": item.strip(), "shade": str(shade).strip(),
            "stock_qty": 0, "ordered_qty": 0,
        })

    for s in stock:
        if not str(s.get("item", "")).strip():
            continue
        e = entry(s.get("item", ""), s.get("shade", ""), s.get("group", ""))
        e["stock_qty"] += s.get("quantity") or 0

    for o in orders:
        if not str(o.get("item", "")).strip():
            continue
        e = entry(o.get("item", ""), o.get("shade", ""), o.get("group", ""))
        e["ordered_qty"] += o.get("qty") or 0

    out = []
    for e in rows.values():
        e["balance"] = e["stock_qty"] - e["ordered_qty"]
        out.append(e)
    out.sort(key=lambda x: (x["group"], x["item"], x["shade"]))
    return {
        "rows": out,
        "total_stock": sum(r["stock_qty"] for r in out),
        "total_ordered": sum(r["ordered_qty"] for r in out),
        "total_balance": sum(r["balance"] for r in out),
        "short_lines": sum(1 for r in out if r["balance"] < 0),
    }


# ---------- Dashboard ----------
@api_router.get("/stats/dashboard")
async def dashboard_stats():
    rows = await db.order_rows.find().to_list(5000)
    stock = await db.stock_rows.find().to_list(5000)

    status_counts = {"not_ready": 0, "arrived": 0, "ready": 0}
    party_wise = {}
    total_qty = 0
    total_amount = 0
    billed = 0
    for r in rows:
        st = r.get("status", "not_ready")
        status_counts[st] = status_counts.get(st, 0) + 1
        qty = r.get("qty") or 0
        total_qty += qty
        total_amount += r.get("amount") or 0
        if str(r.get("bill_no", "")).strip():
            billed += 1
        party = r.get("party_name", "").strip() or "—"
        e = party_wise.setdefault(party, {"party": party, "rows": 0, "qty": 0, "ready": 0, "pending": 0})
        e["rows"] += 1
        e["qty"] += qty
        if st == "ready":
            e["ready"] += 1
        else:
            e["pending"] += 1

    item_wise = {}
    for r in rows:
        item = r.get("item", "").strip()
        if not item:
            continue
        e = item_wise.setdefault(item, {"item": item, "qty": 0})
        e["qty"] += r.get("qty") or 0

    return {
        "total_rows": len(rows),
        "total_qty": total_qty,
        "total_amount": total_amount,
        "billed_rows": billed,
        "unbilled_rows": len(rows) - billed,
        "status_counts": status_counts,
        "stock_lines": len(stock),
        "stock_qty": sum(s.get("quantity") or 0 for s in stock),
        "parties": await db.parties.count_documents({}),
        "party_wise": sorted(party_wise.values(), key=lambda x: -x["qty"])[:8],
        "item_wise": sorted(item_wise.values(), key=lambda x: -x["qty"])[:8],
    }


def _norm(v):
    return str(v or "").strip().upper()


@api_router.get("/stats/overview")
async def dashboard_overview(conference: str = "", group: str = "", party: str = "", item: str = ""):
    all_orders = await db.order_rows.find().to_list(20000)
    stock = await db.stock_rows.find().to_list(20000)
    sent = await db.company_sent_rows.find().to_list(20000)
    arrived = await db.company_arrived_rows.find().to_list(20000)

    filter_options = {
        "conferences": sorted({str(o.get("conference", "")).strip() for o in all_orders if str(o.get("conference", "")).strip()}),
        "groups": sorted({str(o.get("group", "")).strip() for o in all_orders if str(o.get("group", "")).strip()}),
        "parties": sorted({str(o.get("party_name", "")).strip() for o in all_orders if str(o.get("party_name", "")).strip()}),
        "items": sorted({str(o.get("item", "")).strip() for o in all_orders if str(o.get("item", "")).strip()}),
    }

    def keep_order(o):
        if conference and _norm(o.get("conference")) != _norm(conference):
            return False
        if group and _norm(o.get("group")) != _norm(group):
            return False
        if party and _norm(o.get("party_name")) != _norm(party):
            return False
        if item and _norm(o.get("item")) != _norm(item):
            return False
        return True

    def keep_line(r):
        if group and _norm(r.get("group")) != _norm(group):
            return False
        if item and _norm(r.get("item")) != _norm(item):
            return False
        return True

    applied = {"conference": conference, "group": group, "party": party, "item": item}
    orders = [o for o in all_orders if keep_order(o)]
    stock = [s for s in stock if keep_line(s)]
    sent = [s for s in sent if keep_line(s)]
    arrived = [a for a in arrived if keep_line(a)]

    def key3(g, i, s):
        return (str(g or "").strip().upper(), str(i or "").strip().upper(), str(s or "").strip())

    total_qty = total_amount = billed = 0
    status_counts = {"not_ready": 0, "arrived": 0, "ready": 0}
    party_wise = {}
    item_wise = {}
    for r in orders:
        qty = r.get("qty") or 0
        amt = r.get("amount") or 0
        total_qty += qty
        total_amount += amt
        st = r.get("status", "not_ready")
        status_counts[st] = status_counts.get(st, 0) + 1
        if str(r.get("bill_no", "")).strip():
            billed += 1
        party = str(r.get("party_name", "")).strip() or "—"
        e = party_wise.setdefault(party, {
            "party": party, "rows": 0, "qty": 0, "amount": 0,
            "ready": 0, "pending": 0, "billed": 0, "items": set(),
        })
        e["rows"] += 1
        e["qty"] += qty
        e["amount"] += amt
        if str(r.get("bill_no", "")).strip():
            e["billed"] += 1
        if st == "ready":
            e["ready"] += 1
        else:
            e["pending"] += 1
        item = str(r.get("item", "")).strip()
        if item:
            e["items"].add(item.upper())
            it = item_wise.setdefault(item.upper(), {
                "item": item, "group": str(r.get("group", "")).strip(), "qty": 0, "rows": 0, "parties": set(),
            })
            it["qty"] += qty
            it["rows"] += 1
            if party != "—":
                it["parties"].add(party)

    parties_out = []
    for e in party_wise.values():
        parties_out.append({
            **{k: v for k, v in e.items() if k != "items"},
            "items": len(e["items"]),
            "ready_pct": round(100 * e["ready"] / e["rows"]) if e["rows"] else 0,
        })
    parties_out.sort(key=lambda x: -x["qty"])

    items_out = []
    for e in item_wise.values():
        items_out.append({**{k: v for k, v in e.items() if k != "parties"}, "parties": len(e["parties"])})
    items_out.sort(key=lambda x: -x["qty"])

    # pending with company: sent - arrived
    agg = {}
    for r in sent:
        if not str(r.get("item", "")).strip():
            continue
        e = agg.setdefault(key3(r.get("group"), r.get("item"), r.get("shade")), {
            "group": str(r.get("group", "")).strip(), "item": str(r.get("item", "")).strip(),
            "shade": str(r.get("shade", "")).strip(), "sent_qty": 0, "arrived_qty": 0, "last_date": "",
        })
        e["sent_qty"] += r.get("quantity") or 0
        d = str(r.get("date", "")).strip()
        if d > e["last_date"]:
            e["last_date"] = d
    for r in arrived:
        if not str(r.get("item", "")).strip():
            continue
        e = agg.setdefault(key3(r.get("group"), r.get("item"), r.get("shade")), {
            "group": str(r.get("group", "")).strip(), "item": str(r.get("item", "")).strip(),
            "shade": str(r.get("shade", "")).strip(), "sent_qty": 0, "arrived_qty": 0, "last_date": "",
        })
        e["arrived_qty"] += r.get("quantity") or 0
    company = []
    for e in agg.values():
        company.append({**e, "balance_qty": e["sent_qty"] - e["arrived_qty"]})
    pending_company = sorted([c for c in company if c["balance_qty"] > 0], key=lambda x: -x["balance_qty"])

    shortfalls = [r for r in company_order_rows_sync(orders, stock) if r["to_order"] > 0]
    shortfalls.sort(key=lambda x: -x["to_order"])

    return {
        "kpis": {
            "order_rows": len(orders),
            "order_qty": total_qty,
            "order_amount": total_amount,
            "billed_rows": billed,
            "unbilled_rows": len(orders) - billed,
            "parties": len(parties_out),
            "stock_lines": len(stock),
            "stock_qty": sum(s.get("quantity") or 0 for s in stock),
            "sent_qty": sum(c["sent_qty"] for c in company),
            "arrived_qty": sum(c["arrived_qty"] for c in company),
            "pending_company_qty": sum(c["balance_qty"] for c in pending_company),
            "pending_company_lines": len(pending_company),
            "shortfall_lines": len(shortfalls),
            "shortfall_qty": sum(r["to_order"] for r in shortfalls),
        },
        "status_counts": status_counts,
        "party_wise": parties_out,
        "item_wise": items_out,
        "pending_company": pending_company,
        "shortfalls": shortfalls,
        "filter_options": filter_options,
        "applied": applied,
    }


# ---------- AI assistant (Claude Sonnet 4.6) ----------
class AskIn(BaseModel):
    session_id: str
    message: str


ASSISTANT_SYSTEM = (
    "You are HAANA, the operations analyst for a saree/fabric conference order desk. "
    "You answer strictly from the live DATA SNAPSHOT given to you. "
    "Be short and concrete: numbers first, then a one-line reason or action. "
    "Use plain digits, Indian number style. Write plain text only — never use markdown, asterisks or backticks. "
    "If the snapshot does not contain the answer, say so. "
    "Never invent items, parties or quantities."
)


async def assistant_snapshot():
    ov = await dashboard_overview()
    k = ov["kpis"]
    lines = [
        "KPIS: " + ", ".join(f"{a}={b}" for a, b in k.items()),
        f"STATUS ROWS: {ov['status_counts']}",
        "PARTY WISE (party | rows | qty | ready | pending | billed):",
    ]
    for p in ov["party_wise"][:40]:
        lines.append(f"- {p['party']} | {p['rows']} | {p['qty']} | {p['ready']} | {p['pending']} | {p['billed']}")
    lines.append("TOP ITEMS (item | group | qty | parties):")
    for i in ov["item_wise"][:40]:
        lines.append(f"- {i['item']} | {i['group']} | {i['qty']} | {i['parties']}")
    lines.append("SHORTFALL TO ORDER FROM COMPANY (group | item | shade | demand | stock | to_order):")
    for r in ov["shortfalls"][:60]:
        lines.append(f"- {r['group']} | {r['item']} | {r['shade']} | {r['ordered_qty']} | {r['stock_qty']} | {r['to_order']}")
    lines.append("PENDING AT COMPANY (group | item | shade | sent | arrived | pending | last_sent):")
    for r in ov["pending_company"][:60]:
        lines.append(f"- {r['group']} | {r['item']} | {r['shade']} | {r['sent_qty']} | {r['arrived_qty']} | {r['balance_qty']} | {r['last_date']}")
    return "\n".join(lines)


@api_router.get("/assistant/history/{session_id}")
async def assistant_history(session_id: str):
    docs = await db.assistant_messages.find({"session_id": session_id}).sort("created_at", 1).to_list(200)
    return [{"role": d["role"], "text": d["text"], "created_at": d["created_at"]} for d in docs]


@api_router.post("/assistant/ask")
async def assistant_ask(payload: AskIn):
    from emergentintegrations.llm.chat import LlmChat, UserMessage, TextDelta, StreamDone

    key = os.environ.get("EMERGENT_LLM_KEY")
    if not key:
        raise HTTPException(status_code=500, detail="LLM key not configured")

    question = payload.message.strip()
    if not question:
        raise HTTPException(status_code=400, detail="Message is empty")

    snapshot = await assistant_snapshot()
    history = await db.assistant_messages.find({"session_id": payload.session_id}).sort("created_at", 1).to_list(20)
    transcript = "\n".join(f"{h['role'].upper()}: {h['text']}" for h in history[-10:])

    system_message = ASSISTANT_SYSTEM + "\n\nDATA SNAPSHOT:\n" + snapshot
    if transcript:
        system_message += "\n\nEARLIER IN THIS CONVERSATION:\n" + transcript

    chat = LlmChat(
        api_key=key,
        session_id=payload.session_id,
        system_message=system_message,
    ).with_model("anthropic", "claude-sonnet-4-6")

    await db.assistant_messages.insert_one({
        "session_id": payload.session_id, "role": "user", "text": question, "created_at": now_iso(),
    })

    async def gen():
        full = ""
        try:
            async for ev in chat.stream_message(UserMessage(text=question)):
                if isinstance(ev, TextDelta):
                    full += ev.content
                    yield ev.content
                elif isinstance(ev, StreamDone):
                    break
        except Exception as e:  # surface failure to the user instead of a silent hang
            logging.exception("assistant stream failed")
            yield f"\n[assistant error: {e}]"
        if full:
            await db.assistant_messages.insert_one({
                "session_id": payload.session_id, "role": "assistant", "text": full, "created_at": now_iso(),
            })

    return StreamingResponse(gen(), media_type="text/plain", headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})


@api_router.post("/seed")
async def seed():
    if await db.order_rows.count_documents({}) > 0:
        return {"seeded": False, "message": "Data already present"}
    await db.parties.delete_many({})
    await db.parties.insert_many([
        {"name": "SRI RAM VASTRALAYA, BHELAHI", "page": "36", "phone": "", "city": "", "created_at": now_iso()},
        {"name": "CHAUDHARY VASTRALAYA, DARDHA", "page": "60", "phone": "", "city": "", "created_at": now_iso()},
        {"name": "TRILOK VASTRALAYA, LAKHORA", "page": "37", "phone": "", "city": "", "created_at": now_iso()},
    ])
    orders = [
        ("SRI RAM VASTRALAYA, BHELAHI", "36", "V-URBAN LYCRA", "701", 1, "12.80*1", "AS/58", "ready"),
        ("SRI RAM VASTRALAYA, BHELAHI", "36", "V-URBAN LYCRA", "102", 1, "12.80*1", "AS/58", "ready"),
        ("SRI RAM VASTRALAYA, BHELAHI", "36", "V-URBAN LYCRA", "107", 1, "11.20*1", "", "arrived"),
        ("SRI RAM VASTRALAYA, BHELAHI", "36", "V-FORTUNER", "701", 1, "14.40*1", "AS/58", "ready"),
        ("SRI RAM VASTRALAYA, BHELAHI", "36", "V-WAGON", "102", 1, "", "", "not_ready"),
        ("SRI RAM VASTRALAYA, BHELAHI", "36", "V-WAGON", "108", 1, "9.60*1", "AS/56", "ready"),
        ("SRI RAM VASTRALAYA, BHELAHI", "36", "V-WAGON", "114", 1, "11.20*1", "", "arrived"),
        ("CHAUDHARY VASTRALAYA, DARDHA", "60", "V-LINEN STUDIO", "6305", 1, "8*1", "", "arrived"),
        ("CHAUDHARY VASTRALAYA, DARDHA", "60", "V-LINEN STUDIO", "6325", 1, "", "", "not_ready"),
        ("TRILOK VASTRALAYA, LAKHORA", "37", "V-HARRIER WHITE-1", "701", 1, "16*1", "", "arrived"),
        ("TRILOK VASTRALAYA, LAKHORA", "37", "V-TRIBER -5", "501", 1, "", "", "not_ready"),
    ]
    await db.order_rows.insert_many([
        {
            "party_name": p, "page": pg, "conference": "", "group": "SH ROLL", "item": item, "shade": shade,
            "qty": qty, "mtr": mtr, "rate": 0, "amount": 0, "bill_no": bill,
            "status": st, "row_index": idx, "created_at": now_iso(),
        }
        for idx, (p, pg, item, shade, qty, mtr, bill, st) in enumerate(orders)
    ])
    stock = [
        ("LOYAL PRINT CS", "121"), ("LOYAL PRINT CS", "123"), ("LOYAL PRINT CS", "126"),
        ("CUBA PRINT CS", "1872"), ("CUBA PRINT CS", "1873"), ("CUBA PRINT CS", "1916"),
        ("SENSATIONAL CS", "211"), ("SENSATIONAL CS", "212"),
    ]
    await db.stock_rows.insert_many([
        {"conference": "", "group": "SH ROLL", "item": i, "shade": s, "quantity": 3, "row_index": idx, "created_at": now_iso()}
        for idx, (i, s) in enumerate(stock)
    ])
    return {"seeded": True}


app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=os.environ.get('CORS_ORIGINS', '*').split(','),
    allow_methods=["*"],
    allow_headers=["*"],
)

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)


@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()
